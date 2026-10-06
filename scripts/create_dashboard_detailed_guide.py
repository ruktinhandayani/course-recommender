import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUT_DIR = Path(__file__).resolve().parents[1] / "outputs"
DOCX_PATH = OUT_DIR / "Panduan_Rinci_Per_Halaman_Dashboard_CRE_ML.docx"
PDF_PATH = OUT_DIR / "Panduan_Rinci_Per_Halaman_Dashboard_CRE_ML.pdf"
EXPORT_MANIFEST_PATH = Path(__file__).resolve().parents[1] / "course_recommender_artifacts_final" / "export" / "recommendation_export_manifest.json"

NAVY = "0B2545"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
TEXT = "222222"
MUTED = "5F6B76"
GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"


def active_export_summary() -> str:
    """Describe the export from its current manifest without hardcoding scope."""
    try:
        manifest = json.loads(EXPORT_MANIFEST_PATH.read_text(encoding="utf-8"))
        return (
            f"Active recommendation export saat ini memuat {manifest.get('user_count', 'N/A')} user, "
            f"{manifest.get('recommendation_count', 'N/A')} baris rekomendasi, dan Top-{manifest.get('top_k', 'N/A')}."
        )
    except (OSError, ValueError, TypeError):
        return "Jumlah user dan jumlah baris rekomendasi mengikuti active recommendation export dan dapat diverifikasi melalui recommendation export manifest."


ACTIVE_EXPORT_SUMMARY = active_export_summary()


PAGES = [
    {
        "title": "1. Tampilan Awal dan Mode Penggunaan",
        "purpose": "Halaman awal adalah pintu masuk ke research console. Navigasi ikon di bagian atas mengarah ke Overview, Data Quality, Model Card, Evaluation, Error Analysis, User Search, Export Review, Refresh, dan Raw Response. Urutannya mengikuti alur ilmiah penelitian: mulai dari pemeriksaan data, penjelasan model, pengukuran performa, analisis kesalahan, sampai demonstrasi dan operasi hasil rekomendasi.",
        "shown": [
            ("Navigasi modul", "Setiap ikon mewakili halaman kerja. Tujuannya bukan untuk pengguna LMS umum, melainkan untuk memudahkan peneliti atau penguji menelusuri bukti proses machine learning secara sistematis."),
            ("Mode Test", "Mode Test memakai fixture atau contoh export yang aman untuk menguji tampilan, navigasi, dan kesesuaian respons API. Mode ini berguna ketika ingin memastikan antarmuka berjalan tanpa melakukan proses batch pada data aktual."),
            ("Mode Normal", "Mode Normal memakai export rekomendasi aktual yang tersedia pada artefak sistem. Mode ini digunakan ketika hasil yang diperiksa harus merepresentasikan output model yang sesungguhnya."),
        ],
        "reading": "Penjelasan yang dapat disampaikan kepada dosen adalah: mode Test dan Normal tidak berarti terdapat dua model yang berbeda. Perbedaannya berada pada sumber data yang ditampilkan untuk kebutuhan pengujian antarmuka. Evaluasi model tetap merujuk pada artefak eksperimen yang sama, sehingga angka performanya tidak berubah hanya karena mode tampilan diganti.",
        "research": "Halaman awal menunjukkan bahwa penelitian tidak berhenti pada notebook pelatihan. Hasil model disajikan dalam alur yang dapat diaudit, sehingga meningkatkan keterlacakan, kemudahan demonstrasi, dan kesiapan sistem untuk diuji ulang.",
    },
    {
        "title": "2. Overview",
        "purpose": "Overview adalah ringkasan kondisi layanan dan hasil utama penelitian. Halaman ini dipakai untuk memperoleh gambaran cepat sebelum masuk ke penjelasan teknis pada halaman lain.",
        "shown": [
            ("API Status", "Menunjukkan apakah layanan backend dapat menerima permintaan. Status aktif berarti dashboard terhubung dengan service machine learning; apabila gagal, halaman lain yang mengambil data dari API juga dapat terpengaruh."),
            ("Model Version dan Artifact Check", "Model Version menunjukkan artefak model yang sedang digunakan. Artifact Check memeriksa keberadaan dan konsistensi model, encoder, katalog course, konfigurasi, serta tabel pendukung. Bagian ini membantu memastikan output tidak berasal dari file yang tidak lengkap atau versi yang berbeda."),
            ("Recommendations", f"Menunjukkan jumlah rekomendasi hasil export. {ACTIVE_EXPORT_SUMMARY} Hasil dapat diperiksa tanpa melatih model kembali."),
            ("Metric cards", "Kartu metrik merangkum hasil final Softmax Ranking: Hit@1 58,65%, Hit@5 83,14%, Hit@10 91,60%, MRR@10 68,90%, NDCG@10 74,34%, dan Coverage@10 100,00%."),
        ],
        "reading": "Hit@K dan Recall@K menjawab apakah course target berhasil masuk ke daftar Top-K. Dalam evaluasi satu target per sampel, kedua nilai bisa sama. MRR@K dan NDCG@K memperhatikan posisi target; semakin tinggi nilainya, semakin sering course relevan muncul pada peringkat awal. Coverage menunjukkan luasnya variasi course yang dapat diekspos model, bukan hanya kemampuan menemukan target pada sampel tertentu.",
        "research": "Overview cocok digunakan sebagai pembuka presentasi: sistem aktif, artefak tersedia, hasil export tersedia, dan performa utama dapat ditunjukkan dalam satu halaman. Setelah itu, tiap klaim dapat ditelusuri ke halaman data, model, dan evaluasi.",
    },
    {
        "title": "3. Data Quality",
        "purpose": "Data Quality menjelaskan bagaimana data interaksi dipersiapkan sebelum masuk ke model. Halaman ini penting karena kualitas dan pembagian data menentukan validitas hasil evaluasi rekomendasi.",
        "shown": [
            ("Temporal Split Audit", "Bagian ini memisahkan Train Core, Validation Target, Pretest History, dan Test Target. Urutan waktu dipertahankan agar model belajar dari interaksi masa lalu kemudian diuji untuk memprediksi course berikutnya pada periode sesudahnya."),
            ("Noise Cleaning", "Menampilkan proses penyaringan course dummy, sandbox, test, atau entri noise. Pembersihan dilakukan agar pola yang dipelajari model merepresentasikan aktivitas pembelajaran yang relevan, bukan aktivitas teknis atau data percobaan."),
            ("Interaction History", "Menunjukkan sebaran panjang riwayat course per user. Informasi ini berguna untuk mengetahui apakah mayoritas user memiliki riwayat yang cukup untuk membentuk sequence, serta untuk memahami kemungkinan munculnya fallback pada inferensi."),
            ("Course Popularity", "Memperlihatkan frekuensi course pada data train. Hal ini penting untuk membaca apakah rekomendasi mungkin dipengaruhi course yang sangat populer, sehingga evaluasi tidak hanya dilihat dari akurasi tetapi juga dari keragaman hasil."),
        ],
        "reading": "Poin utama yang perlu dijelaskan adalah temporal split. Jika data masa depan masuk ke data train, model dapat memperoleh informasi yang seharusnya belum tersedia ketika prediksi dilakukan. Kondisi tersebut disebut data leakage dan dapat membuat nilai evaluasi tampak lebih tinggi dari kemampuan sebenarnya. Karena itu, pemisahan berbasis waktu membuat skenario uji lebih mendekati penggunaan nyata.",
        "research": "Halaman ini menjadi bukti metodologis bahwa proses preprocessing, penyaringan noise, dan pembagian train-test tidak dilakukan secara sembarangan. Bagian ini mendukung pembahasan validitas eksperimen pada skripsi atau artikel.",
    },
    {
        "title": "4. Model Card",
        "purpose": "Model Card mendokumentasikan model yang digunakan beserta konfigurasi dan batasannya. Fungsinya adalah membuat keputusan desain model dapat dipahami dan direplikasi oleh pembaca penelitian.",
        "shown": [
            ("Architecture", "Model menggunakan Encoder-Decoder LSTM dengan Attention. Encoder membaca urutan riwayat course user, decoder membentuk representasi user, dan attention memberi bobot lebih besar pada bagian riwayat yang paling relevan terhadap prediksi."),
            ("Strategi ranking", "Softmax Ranking digunakan sebagai urutan rekomendasi utama. Cosine Similarity Ranking tetap ditampilkan sebagai skor alternatif untuk perbandingan, bukan sebagai primary ranking."),
            ("Input and Vocabulary", "Model menerima maksimal 10 interaksi terbaru sebagai input sequence. Seluruh riwayat course yang tersedia sebelum prediksi digunakan untuk seen-course masking, sehingga course yang telah diikuti dikeluarkan dari kandidat rekomendasi."),
            ("Representation", "Menampilkan ukuran embedding dan latent dimension. Parameter ini menjelaskan kapasitas representasi model dalam menangkap hubungan antar-course dan pola riwayat user."),
            ("Reproducibility", "Memuat random seed, threshold course pada data train, repeated-run seeds, serta protokol evaluasi. Catatan ini diperlukan agar pelatihan dan pengukuran performa dapat diulang dengan kondisi yang konsisten."),
        ],
        "reading": "Attention tidak berarti model langsung mengetahui alasan bisnis di balik rekomendasi. Attention membantu memberi penekanan pada elemen sequence yang lebih relevan secara komputasional. Karena itu, hasilnya tetap perlu dibaca bersama metadata course, level, dan halaman Error Analysis.",
        "research": "Model Card meningkatkan transparansi penelitian. Dosen dapat melihat bukan hanya nama algoritma, tetapi juga cara model menerima input, menghasilkan ranking, dan bagaimana eksperimen dijaga agar dapat direproduksi.",
    },
    {
        "title": "5. Evaluation",
        "purpose": "Evaluation menampilkan pengukuran performa model pada data test temporal. Halaman ini adalah dasar utama untuk menyatakan seberapa baik model merekomendasikan course yang relevan.",
        "shown": [
            ("Metric Snapshot", "Menampilkan ringkasan metrik utama agar pembaca dapat langsung melihat hasil model tanpa membuka seluruh tabel. Angka yang ditampilkan harus dibaca bersama nilai K dan metode ranking yang tercantum pada kartu."),
            ("Top-K Performance", "Menyajikan hasil pada beberapa nilai K. K adalah jumlah rekomendasi teratas yang diperbolehkan. Contohnya, K=5 berarti sistem dinilai dari lima course teratas yang direkomendasikan."),
            ("Coverage", "Mengevaluasi seberapa luas katalog course muncul dalam hasil rekomendasi. Nilai tinggi menunjukkan model tidak hanya mengekspos sedikit course tertentu."),
            ("Strict Comparison", "Menampilkan robustness check dengan strict_train_history. Default evaluation menggunakan full_pre_target_history. Hasil strict adalah pemeriksaan ketahanan tambahan, bukan bukti bahwa default evaluation mengalami data leakage."),
        ],
        "reading": "Pada final temporal test, Softmax Ranking menghasilkan Hit@1 58,65%, Hit@5 83,14%, Hit@10 91,60%, MRR@10 68,90%, dan NDCG@10 74,34%. Nilai tersebut mengukur apakah course target berikutnya muncul pada ranking offline temporal test; nilai ini tidak berarti persentase user yang pasti puas.",
        "research": "Halaman Evaluation mendukung klaim empiris penelitian. Penyampaian yang baik kepada dosen adalah membedakan antara kemampuan model pada data uji dan dampak nyata di organisasi; validasi pengguna atau eksperimen online diperlukan apabila ingin mengukur kepuasan atau dampak bisnis secara langsung.",
    },
    {
        "title": "6. Error Analysis",
        "purpose": "Error Analysis tidak hanya menampilkan nilai keberhasilan, tetapi juga menjelaskan pola ketika target tidak berada dalam rekomendasi. Halaman ini membantu pembahasan keterbatasan dan peluang perbaikan model.",
        "shown": [
            ("Perbandingan ranking", "Menampilkan perbandingan Similarity Ranking dan Softmax Ranking melalui Hit@10, Miss@10, dan total sampel. Tujuannya adalah melihat pendekatan mana yang lebih sering memasukkan target ke Top-10."),
            ("Panjang riwayat user", "Menunjukkan hubungan antara kecukupan sequence dan hasil prediksi. User dengan riwayat yang lebih pendek biasanya menyediakan konteks lebih sedikit bagi model."),
            ("Frekuensi target", "Menunjukkan apakah course target sering atau jarang muncul dalam data train. Target yang jarang dapat lebih sulit dipelajari karena jumlah contoh pelatihannya terbatas."),
            ("Rank Outcome Distribution", "Membagi hasil menjadi target pada Rank 1, Rank 2-5, Rank 6-10, atau tidak masuk Top-10. Distribusi ini lebih informatif daripada sekadar satu angka akurasi karena menunjukkan seberapa dekat model dengan target."),
        ],
        "reading": "Jika target berada pada Rank 6-10, model sebenarnya masih mengenali course tersebut, tetapi kualitas urutannya belum optimal. Jika target tidak masuk Top-10, penyebabnya dapat berasal dari riwayat yang terlalu pendek, pola user yang belum cukup terwakili, atau course target yang langka. Penjelasan seperti ini membuat pembahasan hasil lebih kritis dan tidak hanya berfokus pada angka tinggi.",
        "research": "Bagi skripsi, Error Analysis dapat menjadi dasar rekomendasi pengembangan, misalnya penanganan cold-start, pengayaan fitur kompetensi, penambahan data interaksi, atau eksperimen strategi ranking yang lain.",
    },
    {
        "title": "7. User Search",
        "purpose": "User Search digunakan untuk menampilkan dan memeriksa hasil rekomendasi individual pengguna dari active recommendation export yang telah dihasilkan oleh pipeline inferensi batch.",
        "shown": [
            ("Input User ID", "Peneliti memasukkan ID user yang tersedia pada active recommendation export. User ID diperlakukan sebagai teks exact sehingga leading zero tidak dihapus atau ditebak."),
            ("Daftar rekomendasi", "Hasil ditampilkan secara terurut. Rank menunjukkan posisi course; Score adalah skor hasil model; Similarity dan Softmax membantu melihat komponen penilaian; sedangkan course, kompetensi, dan level memberi konteks semantik bagi reviewer."),
            ("Reason", "Menjelaskan status hasil. Pada kondisi normal, reason mendukung pembacaan rekomendasi. Pada kondisi tertentu, reason dapat menunjukkan fallback seperti history_too_short atau user_history_not_found."),
        ],
        "reading": "Halaman ini sebaiknya didemonstrasikan dengan memilih satu User ID yang tersedia dalam active export. Jelaskan bahwa halaman membaca batch result yang sudah dihasilkan, lalu tampilkan rank, Softmax score utama, Similarity score alternatif, dan reason agar output dapat diperiksa secara masuk akal.",
        "research": "User Search berperan sebagai bukti audit operasional terhadap active recommendation export. Walaupun bukan penjelasan kausal penuh, skor, urutan, dan kondisi fallback dapat diamati oleh penguji.",
    },
    {
        "title": "8. Export Review",
        "purpose": "Export Review menyediakan pemeriksaan rekomendasi dalam skala batch. Berbeda dengan User Search yang fokus pada satu user, halaman ini menunjukkan hasil yang telah disiapkan untuk seluruh populasi user yang diproses.",
        "shown": [
            ("Tabel export", "Setiap baris umumnya memuat User ID, urutan ranking, course yang direkomendasikan, score, kompetensi, level, dan reason. Struktur ini menjaga output tetap dapat ditelusuri dari user sampai course yang disarankan."),
            ("Pagination", "Data export dapat berjumlah sangat besar, sehingga ditampilkan per halaman. Pagination memungkinkan reviewer berpindah antarbagian data tanpa membebani antarmuka dan tanpa kehilangan struktur kolom."),
            ("Konsistensi output", "Reviewer dapat memeriksa apakah setiap user memperoleh jumlah Top-K yang diharapkan, apakah ranking tersusun benar, dan apakah metadata course tetap tersedia pada hasil export."),
        ],
        "reading": "Saat menjelaskan kepada dosen, tekankan bahwa export batch adalah jembatan dari model ke penggunaan berikutnya. Dengan output terstruktur, rekomendasi dapat ditinjau, disimpan, atau diintegrasikan ke sistem lain tanpa perlu meminta model melakukan inferensi ulang untuk setiap halaman.",
        "research": "Halaman ini memperlihatkan skalabilitas operasional dari hasil penelitian: model bukan hanya menghasilkan metrik, melainkan juga menghasilkan data rekomendasi yang siap ditinjau dan dipakai pada level banyak user.",
    },
    {
        "title": "9. Refresh dan Raw Response",
        "purpose": "Halaman Refresh mengelola pembentukan ulang export rekomendasi, sedangkan Raw Response membuka respons API yang menjadi sumber data dashboard. Keduanya mendukung pengoperasian dan verifikasi sistem.",
        "shown": [
            ("Top-K", "Menentukan jumlah course rekomendasi yang disimpan untuk setiap user pada proses export baru. Nilai yang lebih besar memberi pilihan lebih banyak, tetapi juga memperbesar ukuran hasil."),
            ("Max Users", "Membatasi jumlah user yang diproses. Parameter ini berguna untuk pengujian terbatas sebelum proses dijalankan pada seluruh user."),
            ("Refresh log dan status", "Menampilkan tahapan atau hasil proses batch. Bagian ini membantu memastikan proses berhasil, memantau kegagalan, dan mencatat output yang dihasilkan."),
            ("Raw Response", "Menampilkan payload JSON untuk endpoint seperti health, user, export, evaluation, dan refresh. Halaman ini terutama ditujukan untuk verifikasi kontrak API dan debugging, bukan untuk pembacaan hasil oleh pengguna umum."),
        ],
        "reading": "Refresh tidak sama dengan melatih model ulang. Refresh menggunakan model dan artefak yang sudah tersedia untuk menghasilkan kembali export rekomendasi. Jika model atau data latih ingin berubah, proses pelatihan dan validasi eksperimen harus dilakukan terlebih dahulu. Pemisahan ini menjaga proses operasional tidak mengubah hasil penelitian secara tidak terkontrol.",
        "research": "Adanya Refresh dan Raw Response menunjukkan bahwa sistem memiliki jejak operasional yang jelas. Dashboard dapat memperlihatkan dari mana data berasal, bagaimana output dihasilkan, dan bagaimana kegagalan layanan dapat ditelusuri.",
    },
]


def set_font(run, size=None, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def cell_margin(cell, value=120):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)


def style_table(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
        column.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_font(run, 9, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_labelled(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    lead = p.add_run(label + ". ")
    set_font(lead, 10.6, BLUE_DARK, True)
    body = p.add_run(text)
    set_font(body, 10.6, TEXT)


def add_paragraph(doc, text, title=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if title:
        lead = p.add_run(title + ". ")
        set_font(lead, 10.8, NAVY, True)
    run = p.add_run(text)
    set_font(run, 10.8, TEXT)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    style_table(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, PALE_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run, 10.2, NAVY)


def doc_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, 17, BLUE, True)


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.8)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    header = section.header.paragraphs[0]
    header.text = "CRE ML Recommendation Service | Panduan Rinci Interpretasi Dashboard"
    for run in header.runs:
        set_font(run, 8.8, MUTED)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("Panduan Rinci per Halaman"), 24, NAVY, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    set_font(subtitle.add_run("Dashboard Course Recommendation for Employees (CRE) - Machine Learning Research Console"), 12, MUTED)

    meta = doc.add_table(rows=3, cols=2)
    style_table(meta, [2200, 7160])
    rows = [
        ("Tujuan", "Menjelaskan fungsi, isi, cara membaca, dan relevansi penelitian pada setiap halaman dashboard."),
        ("Sasaran pembaca", "Mahasiswa sebagai presenter, dosen pembimbing/penguji, dan reviewer yang ingin menelusuri proses machine learning."),
        ("Catatan", "Dashboard merupakan research console untuk validasi dan demonstrasi model, bukan antarmuka LMS pegawai."),
    ]
    for row, (key, value) in zip(meta.rows, rows):
        shade(row.cells[0], GRAY)
        for cell, text, bold in ((row.cells[0], key, True), (row.cells[1], value, False)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(text), 10, NAVY if bold else TEXT, bold)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(10)
    intro.paragraph_format.space_after = Pt(0)
    intro.paragraph_format.line_spacing = 1.1
    set_font(intro.add_run("Cara menggunakan panduan ini: setiap bagian dapat disampaikan berurutan saat presentasi. Mulailah dari tujuan halaman, jelaskan elemen yang muncul, kemudian tutup dengan cara membaca dan relevansinya terhadap penelitian."), 10.5, TEXT)

    for page in PAGES:
        pagebreak = doc.add_paragraph()
        pagebreak.add_run().add_break(WD_BREAK.PAGE)
        doc_heading(doc, page["title"])
        add_paragraph(doc, page["purpose"], "Fungsi halaman")
        sub = doc.add_paragraph()
        sub.paragraph_format.space_before = Pt(3)
        sub.paragraph_format.space_after = Pt(4)
        set_font(sub.add_run("Bagian yang ditampilkan"), 12, BLUE_DARK, True)
        for label, text in page["shown"]:
            add_labelled(doc, label, text)
        add_callout(doc, "Cara membaca dan menjelaskan kepada dosen: " + page["reading"])
        follow = doc.add_paragraph()
        follow.paragraph_format.space_before = Pt(8)
        follow.paragraph_format.space_after = Pt(0)
        follow.paragraph_format.line_spacing = 1.1
        lead = follow.add_run("Relevansi penelitian. ")
        set_font(lead, 10.6, NAVY, True)
        set_font(follow.add_run(page["research"]), 10.6, TEXT)

    doc.core_properties.title = "Panduan Rinci per Halaman Dashboard CRE ML"
    doc.core_properties.author = "CRE ML Research"
    doc.core_properties.subject = "Penjelasan dashboard penelitian recommender system"
    doc.save(DOCX_PATH)


def pdf_header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#5F6B76"))
    canvas.setFont("Helvetica", 8.5)
    canvas.drawString(0.75 * inch, 10.55 * inch, "CRE ML Recommendation Service | Panduan Rinci Interpretasi Dashboard")
    canvas.drawRightString(7.75 * inch, 0.55 * inch, f"Halaman {doc.page}")
    canvas.restoreState()


def build_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#0B2545"), spaceAfter=4, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="SubtitleCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=11.5, leading=14, textColor=colors.HexColor("#5F6B76"), spaceAfter=13))
    styles.add(ParagraphStyle(name="PageHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#2E74B5"), spaceAfter=8))
    styles.add(ParagraphStyle(name="SectionHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.HexColor("#1F4D78"), spaceBefore=4, spaceAfter=3))
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.15, leading=13.0, textColor=colors.HexColor("#222222"), spaceAfter=4))
    styles.add(ParagraphStyle(name="SmallCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.1, leading=11.4, textColor=colors.HexColor("#222222")))
    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch, topMargin=0.95 * inch, bottomMargin=0.8 * inch, title="Panduan Rinci per Halaman Dashboard CRE ML", author="CRE ML Research")
    p = lambda text, style="BodyCustom": Paragraph(text, styles[style])
    story = [p("Panduan Rinci per Halaman", "TitleCustom"), p("Dashboard Course Recommendation for Employees (CRE) - Machine Learning Research Console", "SubtitleCustom")]
    metadata = [
        [p("<b>Tujuan</b>", "SmallCustom"), p("Menjelaskan fungsi, isi, cara membaca, dan relevansi penelitian pada setiap halaman dashboard.", "SmallCustom")],
        [p("<b>Sasaran pembaca</b>", "SmallCustom"), p("Mahasiswa sebagai presenter, dosen pembimbing/penguji, dan reviewer penelitian.", "SmallCustom")],
        [p("<b>Catatan</b>", "SmallCustom"), p("Research console untuk validasi dan demonstrasi model, bukan antarmuka LMS pegawai.", "SmallCustom")],
    ]
    table = Table(metadata, colWidths=[1.5 * inch, 5.5 * inch])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7")), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DCE4")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story += [table, Spacer(1, 9), p("<b>Cara menggunakan panduan.</b> Setiap bagian dapat disampaikan berurutan saat presentasi: mulai dari fungsi halaman, elemen yang muncul, cara membacanya, lalu relevansinya terhadap penelitian.")]

    for page in PAGES:
        story.append(PageBreak())
        flow = [p(page["title"], "PageHeading"), p("<b>Fungsi halaman.</b> " + page["purpose"])]
        flow.append(p("Bagian yang ditampilkan", "SectionHeading"))
        for label, text in page["shown"]:
            flow.append(p(f"<b>{label}.</b> {text}"))
        callout = Table([[p("<b>Cara membaca dan menjelaskan kepada dosen.</b> " + page["reading"], "SmallCustom")]], colWidths=[7.0 * inch])
        callout.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E8EEF5")), ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C9DD")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
        flow += [Spacer(1, 3), callout, Spacer(1, 7), p("<b>Relevansi penelitian.</b> " + page["research"])]
        story.append(KeepTogether(flow))
    doc.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
