from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUT_DIR = Path(__file__).resolve().parents[1] / "outputs"
OUT_FILE = OUT_DIR / "Panduan_Interpretasi_Dashboard_CRE_ML.docx"
PDF_FILE = OUT_DIR / "Panduan_Interpretasi_Dashboard_CRE_ML.pdf"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
GRAY_FILL = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
MUTED = "5F6B76"


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
        grid_col.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_font(run, 9, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_font(run, 11, "222222")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13, BLUE if level == 1 else DARK_BLUE, bold=True)
    return p


def add_bullet(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    first = p.add_run(f"{label}. ")
    set_font(first, 11, NAVY, bold=True)
    rest = p.add_run(text)
    set_font(rest, 11, "222222")


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_font(run, 10, DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    run = p2.add_run(text)
    set_font(run, 10.5, "222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.text = "CRE ML Recommendation Service | Panduan Interpretasi Dashboard"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, 9, MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Panduan Interpretasi Dashboard")
    set_font(run, 24, NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Course Recommendation for Employees (CRE) - Machine Learning Research Console")
    set_font(run, 12, MUTED)

    metadata = doc.add_table(rows=3, cols=2)
    set_table_widths(metadata, [2160, 7200])
    metadata_rows = [
        ("Tujuan", "Menjelaskan fungsi halaman dashboard dan cara membaca hasil penelitian."),
        ("Fokus", "Audit data, dokumentasi model, evaluasi ranking, inferensi, export, dan operasi layanan."),
        ("Catatan", "Dashboard ini merupakan alat pendukung penelitian machine learning, bukan tampilan LMS untuk pegawai."),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        set_cell_shading(row.cells[0], GRAY_FILL)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_font(r, 10, NAVY, bold=True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, 10, "222222")

    add_heading(doc, "1. Gambaran Umum Dashboard")
    add_body(doc, "Dashboard CRE ML dirancang untuk memperlihatkan seluruh alur penelitian secara transparan: kualitas data, konfigurasi model, hasil evaluasi, rekomendasi per user, export batch, hingga pembaruan hasil inferensi. Dengan demikian, dashboard membantu membuktikan bahwa model tidak hanya dibangun, tetapi juga dapat diperiksa, dievaluasi, dan digunakan ulang secara terstruktur.")
    add_callout(doc, "Cara membaca mode data", "Mode Test menggunakan fixture export untuk menguji tampilan dan kontrak API. Mode Normal menggunakan export rekomendasi aktual dari artefak model. Metrik evaluasi tetap berasal dari artefak eksperimen yang sama.")

    add_heading(doc, "2. Overview")
    add_body(doc, "Halaman Overview merupakan ringkasan kondisi layanan. Kartu API Status menunjukkan apakah backend aktif; Model Version menunjukkan artefak model yang sedang digunakan; Recommendations menunjukkan jumlah hasil export; sedangkan Artifact Check menunjukkan konsistensi antara model, encoder, katalog course, dan tabel pendukung. Total Dataset Users berjumlah 9,598, Evaluation Samples berjumlah 4,501, sedangkan Active Export Users saat ini hanya 20; ketiga angka ini memiliki arti berbeda dan tidak boleh disamakan.")
    add_body(doc, "Bagian metrik menampilkan ringkasan hasil evaluasi utama. Hit@K dan Recall@K menunjukkan kemampuan menemukan course target dalam Top-K; MRR@K dan NDCG@K menunjukkan kualitas posisi ranking; Coverage menunjukkan luasnya course katalog yang berhasil dijangkau oleh rekomendasi.")

    add_heading(doc, "3. Data Quality")
    add_body(doc, "Data Quality mendokumentasikan tahap preprocessing agar penggunaan data dapat dipertanggungjawabkan. Temporal Split Audit memperlihatkan data Train Core, Validation Target, Pretest History, dan Test Target. Pembagian berdasarkan waktu memastikan model mempelajari interaksi masa lalu untuk memprediksi course berikutnya, sehingga mengurangi risiko data leakage.")
    add_bullet(doc, "Noise Cleaning", "menunjukkan course dummy, sandbox, test, atau noise yang dihapus sebelum pembentukan sequence.")
    add_bullet(doc, "Interaction History", "menunjukkan distribusi jumlah course yang pernah diambil user, sehingga dapat diketahui kecukupan riwayat untuk inferensi.")
    add_bullet(doc, "Course Popularity", "menunjukkan distribusi frekuensi course pada data train dan membantu membaca potensi pengaruh course populer terhadap rekomendasi.")

    add_heading(doc, "4. Model Card")
    add_body(doc, "Model Card adalah dokumentasi teknis dari model aktif. Arsitektur yang digunakan adalah Encoder-Decoder LSTM dengan Attention. Encoder mempelajari urutan riwayat course pengguna, decoder membentuk representasi pengguna, dan attention membantu model memberi bobot lebih besar pada riwayat course yang paling relevan.")
    add_bullet(doc, "Architecture", "menjelaskan jenis model, versi artefak, dan strategi ranking. Softmax Ranking adalah ranking utama; Cosine Similarity Ranking tersedia sebagai skor alternatif.")
    add_bullet(doc, "Input and Vocabulary", "menunjukkan input maksimum 10 interaksi terbaru, sedangkan seluruh history course sebelum waktu prediksi digunakan untuk mengecualikan course yang telah diikuti dari kandidat rekomendasi.")
    add_bullet(doc, "Representation", "menampilkan ukuran embedding dan latent dimension yang digunakan untuk merepresentasikan fitur pembelajaran.")
    add_bullet(doc, "Reproducibility", "menampilkan random seed, threshold course pada data train, repeated-run seeds, dan protokol evaluasi agar eksperimen dapat diulang secara konsisten.")

    add_heading(doc, "5. Evaluation")
    add_body(doc, "Evaluation menampilkan pengukuran performa model pada data test yang dipisahkan secara temporal. Metric Snapshot memberikan tampilan ringkas, sedangkan tabel Top-K Performance, Coverage, dan Strict Comparison menyediakan detail hasil eksperimen.")
    add_body(doc, "Hasil final Softmax Ranking adalah Hit@1 58.65%, Hit@5 83.14%, Hit@10 91.60%, MRR@10 68.90%, NDCG@10 74.34%, dan Coverage@10 100.00%. Default evaluation memakai full_pre_target_history; strict robustness check memakai strict_train_history.")
    eval_table = doc.add_table(rows=1, cols=3)
    set_table_widths(eval_table, [1800, 2520, 5040])
    headers = ["Metrik", "Makna", "Interpretasi"]
    for cell, text in zip(eval_table.rows[0].cells, headers):
        set_cell_shading(cell, GRAY_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, 10, NAVY, bold=True)
    metrics = [
        ("Hit@K / Recall@K", "Target muncul di Top-K", "Menunjukkan kemampuan model menemukan course relevan. Pada satu target per sampel, kedua nilai dapat sama."),
        ("MRR@K", "Posisi target", "Nilai lebih tinggi berarti target lebih sering muncul pada ranking awal."),
        ("NDCG@K", "Kualitas urutan", "Memberi bobot lebih tinggi apabila target berada di peringkat lebih atas."),
        ("Coverage@K", "Jangkauan katalog", "Menunjukkan apakah rekomendasi menjangkau variasi course, tidak hanya course tertentu."),
    ]
    for metric, meaning, interpretation in metrics:
        cells = eval_table.add_row().cells
        for cell, value in zip(cells, [metric, meaning, interpretation]):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_font(r, 9.5, "222222", bold=(cell == cells[0]))

    add_heading(doc, "6. Error Analysis")
    add_body(doc, "Error Analysis menjelaskan pola keberhasilan dan kegagalan rekomendasi. Ringkasan membandingkan Similarity Ranking dan Softmax Ranking melalui Hit@10, Miss@10, total sampel, panjang riwayat user, serta frekuensi target pada data train.")
    add_body(doc, "Rank Outcome Distribution memperlihatkan apakah target berada pada Rank 1, Rank 2-5, Rank 6-10, atau tidak masuk Top-10. Analisis ini penting karena nilai evaluasi saja belum cukup menjelaskan posisi target dan faktor yang menyebabkan miss. Course yang lebih jarang muncul pada data train atau pola riwayat user yang lebih kompleks dapat menjadi sumber keterbatasan model.")

    add_heading(doc, "7. User Search")
    add_body(doc, "User Search menampilkan hasil rekomendasi individual dari active recommendation export, bukan melakukan live/direct inference. Pengguna memasukkan User ID sebagai teks exact; leading zero dipertahankan dan tidak ditebak. Kolom Rank, Course, Softmax, Similarity, Level, dan Reason membantu membaca hasil yang telah diekspor.")
    add_body(doc, "Apabila User ID ada pada dataset sumber tetapi tidak berada pada export aktif, dashboard menampilkan status user_not_in_active_export. Reason lain dapat menunjukkan history_too_short, user_history_not_found, model_inference, atau popularity_fallback sesuai hasil pada export.")

    add_heading(doc, "8. Export Review")
    add_body(doc, "Export Review menampilkan data dari active recommendation export, yang dapat berupa demonstration subset atau full batch export tergantung konfigurasi. Active export saat ini adalah demonstration_subset dengan 20 user, Top-K 10, dan 200 recommendation rows; export ini tidak merepresentasikan seluruh 9,598 pengguna.")
    add_body(doc, "Setiap baris berisi User ID, ranking, course, score, kompetensi, level, dan reason. Fitur pagination digunakan agar export dapat diperiksa secara bertahap tanpa menampilkan seluruh baris sekaligus.")

    add_heading(doc, "9. Refresh dan Raw Response")
    add_body(doc, "Refresh digunakan untuk menghasilkan ulang file export tanpa melatih model dari awal. Parameter Top-K menentukan jumlah rekomendasi per user, sedangkan Max Users mengatur subset bila diisi; export penuh hanya dibuat bila dikonfigurasi secara eksplisit. Log proses dan ringkasan hasil membantu memantau keberhasilan proses batch.")
    add_body(doc, "Raw Response menampilkan payload JSON dari health, user, export, evaluation, dan refresh. Halaman ini mendukung verifikasi struktur API, debugging, dan pembuktian bahwa data frontend berasal dari respons backend yang terdefinisi.")

    add_heading(doc, "10. Kesimpulan")
    add_body(doc, "Secara keseluruhan, dashboard CRE ML berfungsi sebagai bukti operasional dan ilmiah dari penelitian. Data diaudit sebelum digunakan, model didokumentasikan, performa dievaluasi dengan metrik ranking, kesalahan dianalisis, rekomendasi dapat diperiksa per user maupun batch, dan proses refresh dapat dipantau. Struktur ini mendukung transparansi, reproduktibilitas, serta kesiapan sistem untuk dipresentasikan dalam skripsi maupun artikel ilmiah.")

    doc.core_properties.title = "Panduan Interpretasi Dashboard CRE ML"
    doc.core_properties.subject = "Dokumentasi dashboard penelitian machine learning"
    doc.core_properties.author = "CRE ML Research"
    doc.save(OUT_FILE)


def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8.5)
    canvas.setFillColor(colors.HexColor("#5F6B76"))
    canvas.drawString(0.75 * inch, 10.55 * inch, "CRE ML Recommendation Service | Panduan Interpretasi Dashboard")
    canvas.drawRightString(7.75 * inch, 0.55 * inch, f"Halaman {doc.page}")
    canvas.restoreState()


def build_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="DocTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#0B2545"), spaceAfter=4, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="Subtitle", parent=styles["BodyText"], fontName="Helvetica", fontSize=11.5, leading=14, textColor=colors.HexColor("#5F6B76"), spaceAfter=14))
    styles.add(ParagraphStyle(name="H1Custom", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.2, leading=12.8, textColor=colors.HexColor("#222222"), spaceAfter=4))
    styles.add(ParagraphStyle(name="SmallCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.1, leading=11.4, textColor=colors.HexColor("#222222")))
    doc = SimpleDocTemplate(str(PDF_FILE), pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch, topMargin=0.95 * inch, bottomMargin=0.8 * inch, title="Panduan Interpretasi Dashboard CRE ML", author="CRE ML Research")
    story = []
    P = lambda text, style="BodyCustom": Paragraph(text, styles[style])
    story += [P("Panduan Interpretasi Dashboard", "DocTitle"), P("Course Recommendation for Employees (CRE) - Machine Learning Research Console", "Subtitle")]
    meta = [[P("<b>Tujuan</b>", "SmallCustom"), P("Menjelaskan fungsi halaman dashboard dan cara membaca hasil penelitian.", "SmallCustom")], [P("<b>Fokus</b>", "SmallCustom"), P("Audit data, dokumentasi model, evaluasi ranking, inferensi, export, dan operasi layanan.", "SmallCustom")], [P("<b>Catatan</b>", "SmallCustom"), P("Dashboard ini merupakan alat pendukung penelitian machine learning, bukan tampilan LMS untuk pegawai.", "SmallCustom")]]
    meta_table = Table(meta, colWidths=[1.45 * inch, 5.55 * inch])
    meta_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7")), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DCE4")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story += [meta_table, Spacer(1, 8)]

    sections = [
        ("1. Gambaran Umum Dashboard", ["Dashboard CRE ML memperlihatkan alur penelitian secara transparan: kualitas data, konfigurasi model, hasil evaluasi, rekomendasi per user, export batch, dan pembaruan inferensi. Dashboard berfungsi sebagai bukti bahwa model dapat diperiksa, dievaluasi, dan digunakan ulang secara terstruktur.", "<b>Mode data.</b> Mode Test memakai fixture export untuk pengujian tampilan dan kontrak API. Mode Normal memakai export rekomendasi aktual. Metrik evaluasi tetap berasal dari artefak eksperimen yang sama."]),
        ("2. Overview", ["Overview merangkum kondisi layanan. API Status menunjukkan backend aktif, Model Version menunjukkan artefak yang digunakan, Recommendations menunjukkan jumlah hasil export, dan Artifact Check menunjukkan konsistensi model, encoder, katalog course, serta tabel pendukung. Total Dataset Users 9,598, Evaluation Samples 4,501, dan Active Export Users 20 adalah angka dengan arti berbeda.", "Kartu metrik menampilkan Hit@K dan Recall@K untuk kemampuan menemukan target; MRR@K dan NDCG@K untuk kualitas posisi ranking; serta Coverage untuk jangkauan course dalam katalog."]),
        ("3. Data Quality", ["Data Quality mendokumentasikan preprocessing. Temporal Split Audit memperlihatkan Train Core, Validation Target, Pretest History, dan Test Target. Pembagian berbasis waktu memastikan model mempelajari interaksi masa lalu untuk memprediksi course berikutnya dan mengurangi risiko data leakage.", "Noise Cleaning menampilkan course dummy, sandbox, test, atau noise yang dihapus. Interaction History menjelaskan distribusi riwayat course per user. Course Popularity menunjukkan frekuensi course pada data train sebagai konteks kemungkinan bias terhadap course populer."]),
        ("4. Model Card", ["Model Card mendokumentasikan model Encoder-Decoder LSTM dengan Attention. Softmax Ranking adalah ranking utama; Cosine Similarity Ranking adalah skor alternatif. Input model memakai maksimum 10 interaksi terbaru, sedangkan seluruh history sebelum prediksi mengecualikan course yang telah diikuti.", "Architecture menjelaskan strategi ranking; Input and Vocabulary menjelaskan sequence, course, kompetensi, dan level yang dipahami model; Representation menampilkan embedding dan latent dimension; sedangkan Reproducibility memuat random seed, threshold training, repeated-run seeds, dan protokol evaluasi."]),
        ("5. Evaluation", ["Evaluation menampilkan performa model pada data test temporal. Hasil Softmax final: Hit@1 58.65%, Hit@5 83.14%, Hit@10 91.60%, MRR@10 68.90%, NDCG@10 74.34%, Coverage@10 100.00%. Default memakai full_pre_target_history dan strict robustness check memakai strict_train_history."]),
        ("6. Error Analysis", ["Error Analysis membandingkan Similarity Ranking dan Softmax Ranking melalui Hit@10, Miss@10, total sampel, panjang riwayat user, dan frekuensi target pada data train. Rank Outcome Distribution memperlihatkan apakah target muncul pada Rank 1, Rank 2-5, Rank 6-10, atau tidak masuk Top-10.", "Bagian ini membantu menjelaskan keterbatasan model. Misalnya, target yang jarang muncul pada data train atau riwayat user yang lebih kompleks dapat lebih sulit diprediksi."]),
        ("7. User Search", ["User Search menampilkan hasil rekomendasi individual dari active recommendation export, bukan live/direct inference. User ID diperlakukan sebagai exact text sehingga leading zero dipertahankan.", "Jika ID ada pada data sumber tetapi tidak di export aktif, dashboard menampilkan user_not_in_active_export. Reason lain dapat menunjukkan history_too_short, user_history_not_found, model_inference, atau popularity_fallback."]),
        ("8. Export Review", ["Export Review menampilkan active recommendation export, yang dapat berupa demonstration subset atau full batch export tergantung konfigurasi. Active export saat ini adalah demonstration_subset: 20 user, Top-K 10, dan 200 recommendation rows.", "Export aktif tidak merepresentasikan seluruh 9,598 pengguna. Setiap baris memuat User ID, ranking, course, score, kompetensi, level, dan reason; pagination memungkinkan pemeriksaan bertahap."]),
        ("9. Refresh dan Raw Response", ["Refresh menghasilkan ulang export tanpa melatih model dari awal. Top-K menentukan jumlah rekomendasi per user, Max Users digunakan untuk pengujian terbatas, dan log proses memantau keberhasilan batch.", "Raw Response menampilkan payload JSON health, user, export, evaluation, dan refresh untuk verifikasi kontrak API serta debugging."]),
        ("10. Kesimpulan", ["Dashboard CRE ML mendukung transparansi dan reproduktibilitas penelitian: data diaudit, model didokumentasikan, performa dievaluasi, kesalahan dianalisis, rekomendasi dapat diperiksa per user maupun batch, dan pembaruan export dapat dipantau. Struktur ini siap digunakan sebagai pendukung presentasi skripsi dan artikel ilmiah."]),
    ]
    for heading, paragraphs in sections:
        section_flowables = [P(heading, "H1Custom")]
        section_flowables.extend(P(paragraph) for paragraph in paragraphs)
        story.append(KeepTogether(section_flowables))

        if heading == "5. Evaluation":
            table_data = [[P("<b>Metrik</b>", "SmallCustom"), P("<b>Makna</b>", "SmallCustom"), P("<b>Interpretasi</b>", "SmallCustom")], [P("Hit@K / Recall@K", "SmallCustom"), P("Target masuk Top-K", "SmallCustom"), P("Menunjukkan kemampuan menemukan course relevan. Pada satu target per sampel, keduanya dapat sama.", "SmallCustom")], [P("MRR@K", "SmallCustom"), P("Posisi target", "SmallCustom"), P("Nilai lebih tinggi berarti target lebih sering berada di ranking awal.", "SmallCustom")], [P("NDCG@K", "SmallCustom"), P("Kualitas urutan", "SmallCustom"), P("Memberi bobot lebih tinggi untuk target di peringkat atas.", "SmallCustom")], [P("Coverage@K", "SmallCustom"), P("Jangkauan katalog", "SmallCustom"), P("Menunjukkan variasi course yang berhasil diekspos model.", "SmallCustom")]]
            metric_table = Table(table_data, colWidths=[1.35 * inch, 1.75 * inch, 3.9 * inch], repeatRows=1)
            metric_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D6DCE4")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story += [metric_table, Spacer(1, 5)]
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


if __name__ == "__main__":
    build_document()
    build_pdf()
    print(OUT_FILE)
    print(PDF_FILE)
